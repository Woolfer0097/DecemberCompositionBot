import telebot
from telebot import types  # для указание типов
import random
import config
import os
from dirt_tongue import is_dirt
from datetime import datetime
import docx
from docx2pdf import convert
import pythoncom


DATA_DIRECTORY = "./data/"
THEORY_DIRECTORY = "theory/"
SECTION1_NAME = "ДУХОВНО-НРАВСТВЕННЫЕ ОРИЕНТИРЫ В ЖИЗНИ ЧЕЛОВЕКА"
SECTION2_NAME = "СЕМЬЯ, ОБЩЕСТВО, ОТЕЧЕСТВО В ЖИЗНИ ЧЕЛОВЕКА"
SECTION3_NAME = "ПРИРОДА И КУЛЬТУРА В ЖИЗНИ ЧЕЛОВЕКА"
SECTIONS = [SECTION1_NAME, SECTION2_NAME, SECTION3_NAME]
detector = is_dirt()


def read_file(filename):
    with open(DATA_DIRECTORY + filename, "r", encoding="utf-8") as file_read:
        data = file_read.read()
    return data


def get_data(section_number=0):
    if section_number == 0:
        whole_data = []
        for i in range(1, 4):
            whole_data.extend(read_file(f"{i}.txt").split()[1:])
        return {"data": whole_data}
    data = read_file(f"{section_number}.txt").splitlines()
    section_name = data[0]
    return {"section_name": section_name, "data": data[1:]}


def get_random_variant(count=6):
    problems = dict()
    problems["sections"] = 1, 2, 3
    problems["problems"] = []
    standart = [1, 1, 2, 2, 3, 3]
    for i in range(count % 3):
        standart.insert(0, 1)
    for j in [1, 1, 2, 2, 3, 3]:
        data = get_data(j)["data"]
        problem_number = random.randint(0, len(data))
        while problem_number in [prob["problem_number"] for prob in problems["problems"]]:
            problem_number = random.randint(0, len(data))
        problems["problems"].append({"problem_number": problem_number,
                                     "section": j,
                                     "problem": data[problem_number]})
    return problems


bot = telebot.TeleBot(config._token)

menu_btn = types.KeyboardButton("Меню")

work_grid = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
btn1 = types.KeyboardButton("Команды")
btn2 = types.KeyboardButton("Реальный вариант")
btn3 = types.KeyboardButton("Теория")
btn4 = types.KeyboardButton("Раздел")
work_grid.add(btn1, btn2, btn3, btn4, menu_btn)

menu = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=1)
btn_section1 = types.KeyboardButton("Раздел ДУХОВНО-НРАВСТВЕННЫЕ ОРИЕНТИРЫ В ЖИЗНИ ЧЕЛОВЕКА 1")
btn_section2 = types.KeyboardButton("Раздел СЕМЬЯ, ОБЩЕСТВО, ОТЕЧЕСТВО В ЖИЗНИ ЧЕЛОВЕКА 2")
btn_section3 = types.KeyboardButton("Раздел ПРИРОДА И КУЛЬТУРА В ЖИЗНИ ЧЕЛОВЕКА 3")
menu.add(btn_section1, btn_section2, btn_section3, menu_btn)

theory_menu = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
btn_1 = types.KeyboardButton("Произведения")
btn_2 = types.KeyboardButton("Структура")
theory_menu.add(btn_1, btn_2, menu_btn)


@bot.message_handler(commands=['start'])
def start(message):
    bot.send_message(message.chat.id,
                     text=f"Привет, {message.from_user.first_name}!\n{read_file('greeting.txt')}",
                     reply_markup=work_grid)


def commands(message):
    bot.send_message(message.chat.id, read_file("commands.txt"), reply_markup=work_grid)


def real(message):
    bot.send_message(message.chat.id, "Обработка запроса...")
    count = 6
    is_unusual_variant = message.text.isdigit()
    if is_unusual_variant:
        count = int(message.text.split()[-1])
    output_message = []
    doc_data = []
    variant_data = get_random_variant(count)

    for section in variant_data["sections"]:
        for problem in variant_data["problems"]:
            if problem["section"] == section:
                problem_data = problem['problem'].split()
                theme_number = problem_data[0]
                problem_content = " ".join(problem_data[1:]).lstrip()
                output_message.append(f"{problem['section']}. {problem['problem']}\n")
                doc_data.append([str(theme_number), str(problem_content)])

    # Создание PDF-файла
    if not is_unusual_variant:
        pythoncom.CoInitializeEx(0)
        doc = docx.Document(DATA_DIRECTORY + "Blank.docx")
        table = doc.tables[0]
        for i in range(len(output_message)):
            table.cell(i + 1, 0).text = doc_data[i][0]
            table.cell(i + 1, 1).text = doc_data[i][1]

        doc.save(DATA_DIRECTORY + f"tmp/Blank_temp.docx")

        convert(DATA_DIRECTORY + "tmp/Blank_temp.docx", DATA_DIRECTORY + "tmp/Вариант для печати.pdf")
        bot.send_message(message.chat.id, "\n".join(output_message))
        with open(DATA_DIRECTORY + "tmp/Вариант для печати.pdf", "rb") as file:
            bot.send_document(message.chat.id, document=file, caption="Вариант для печати")
        file.close()
        os.remove(DATA_DIRECTORY + f"tmp/Blank_temp.docx")
        os.remove(DATA_DIRECTORY + "tmp/Вариант для печати.pdf")
    else:
        bot.send_message(message.chat.id, "\n".join(output_message))


def section(message, text):
    section_number = text.split()[-1]
    if not section_number.isdigit():
        bot.send_message(message.chat.id, "Выберите раздел в меню", reply_markup=menu)
    else:
        data = get_data(section_number)
        bot.send_message(message.chat.id, data["section_name"])
        for index, i in enumerate(range(30, len(data["data"]), 30)):
            bot.send_message(message.chat.id, "\n".join(data["data"][index * 30 + 1:i]), reply_markup=work_grid)


def theory(message):
    bot.send_message(message.chat.id, "Выберите, что хотите изучить", reply_markup=theory_menu)


@bot.message_handler(content_types=['text'])
def get_text_messages(message):
    text = message.text.lower()
    if "раздел" in text:
        section(message, text)
    elif "реальный вариант" in text:
        real(message)
    elif text in ["команды", "команда", "rjvfyls", "rjvfylf"]:
        commands(message)
    elif detector(text):
        bot.send_photo(message.chat.id, "https://i.pinimg.com/originals/19/14/65/191465a96c23fb43347e5bad7327645b.jpg",
                       reply_markup=work_grid)
    elif text == "теория":
        theory(message)
    elif text == "произведения":
        with open(DATA_DIRECTORY + THEORY_DIRECTORY + "literature.pdf", "rb") as literature:
            bot.send_document(message.chat.id, literature)
    elif text == "структура":
        with open(DATA_DIRECTORY + THEORY_DIRECTORY + "structure.pdf", "rb") as structure:
            bot.send_document(message.chat.id, structure)
    elif text == "меню":
        bot.send_message(message.chat.id, "Меню", reply_markup=work_grid)
    else:
        markup = types.InlineKeyboardMarkup()
        button_blog = types.InlineKeyboardButton("Мой блог", url='https://t.me/WoolferBlog/18')
        markup.add(button_blog)
        bot.send_photo(message.chat.id,
                       "https://tlgrm.ru/_/stickers/d21/e99/d21e9940-fc86-49ba-9d91-b20a71136040/11.jpg",
                       "Такой команды пока нет!",
                       reply_markup=work_grid)
    print(datetime.now(), message.chat.id, message.from_user.username, message.text, sep=" | ")
    # bot.delete_message(message.chat.id, message.message_id)


def run():
    bot.polling(none_stop=True)


if __name__ == '__main__':
    run()
